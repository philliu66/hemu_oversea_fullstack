"""全局工具：文件解析（Excel 等）。"""

from __future__ import annotations

import json
import re
from io import BytesIO

from fastapi import UploadFile
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from app.core.errors import bad_request
from app.schemas.utils import FileExtractPublic, FileExtractSkippedRow

_COL_LETTER_RE = re.compile(r"^[A-Za-z]{1,3}$")
_FIELD_NAME_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_ALLOWED_EXCEL_SUFFIXES = (".xlsx", ".xlsm")


def _normalize_col_letter(value: str, *, field_name: str) -> str:
    text = str(value or "").strip().upper()
    if not text or not _COL_LETTER_RE.fullmatch(text):
        raise bad_request(f"{field_name} 须为 Excel 列字母，如 A / B / AA")
    try:
        column_index_from_string(text)
    except ValueError as exc:
        raise bad_request(f"{field_name} 不是有效的列字母：{text}") from exc
    return text


def _normalize_field_name(value: object) -> str:
    text = str(value or "").strip()
    if not text or not _FIELD_NAME_RE.fullmatch(text):
        raise bad_request(
            f"extract_field_info 键名无效：{value!r}（须为字母/数字/下划线，且不以数字开头）"
        )
    return text


def parse_extract_field_info(raw: str | dict[str, str]) -> dict[str, str]:
    """
    解析提取字段映射。

    支持：
    - dict：{"itemCode": "B", "grossWeight": "C"}（商品更新请用官网字段名）
    - JSON 字符串：同上
    """
    if isinstance(raw, dict):
        data = raw
    else:
        text = str(raw or "").strip()
        if not text:
            raise bad_request("extract_field_info 不能为空")
        try:
            data = json.loads(text)
        except json.JSONDecodeError as exc:
            raise bad_request(
                "extract_field_info 须为合法 JSON 对象，如 "
                '{"itemCode":"B","grossWeight":"C","baseQt":"D","stepQt":"E","stepWeight":"F"}'
            ) from exc

    if not isinstance(data, dict) or not data:
        raise bad_request("extract_field_info 须为非空 JSON 对象")

    mapping: dict[str, str] = {}
    for key, col in data.items():
        field = _normalize_field_name(key)
        if field in mapping:
            raise bad_request(f"extract_field_info 键名重复：{field}")
        letter = _normalize_col_letter(str(col), field_name=f"extract_field_info.{field}")
        mapping[field] = letter
    return mapping


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        text = format(value, "f").rstrip("0").rstrip(".")
        return text
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def _extract_rows_from_sheet(
    sheet: Worksheet,
    *,
    sheet_index: int,
    field_cols: dict[str, int],
    start_row: int,
    end_row: int | None,
) -> tuple[list[dict[str, str]], list[FileExtractSkippedRow]]:
    max_col = max(field_cols.values()) if field_cols else 1
    field_names = list(field_cols.keys())
    items: list[dict[str, str]] = []
    skipped: list[FileExtractSkippedRow] = []
    last_content_row: int | None = None

    iter_kwargs: dict[str, int] = {
        "min_row": start_row,
        "max_col": max_col,
    }
    if end_row is not None:
        iter_kwargs["max_row"] = end_row

    for row_no, row in enumerate(
        sheet.iter_rows(**iter_kwargs, values_only=True),
        start=start_row,
    ):
        values = row if isinstance(row, tuple) else tuple(row or ())
        record: dict[str, str] = {}
        missing: list[str] = []
        for field, col_idx in field_cols.items():
            cell = values[col_idx - 1] if len(values) >= col_idx else None
            text = _cell_to_str(cell)
            if not text:
                missing.append(field)
            else:
                record[field] = text

        if not missing:
            items.append(record)
            last_content_row = row_no
            continue

        if not record:
            # 映射字段全空：记为空行（表尾连续空行稍后裁掉）
            skipped.append(
                FileExtractSkippedRow(
                    sheet_index=sheet_index,
                    row_index=row_no,
                    reason="empty_row",
                    missing_fields=field_names,
                )
            )
            continue

        # 部分字段有值、部分为空：数据不完整
        last_content_row = row_no
        skipped.append(
            FileExtractSkippedRow(
                sheet_index=sheet_index,
                row_index=row_no,
                reason="missing_fields",
                missing_fields=missing,
            )
        )

    # 裁掉每个 sheet 末尾连续 empty_row，避免 Excel 使用区域尾部空白刷屏
    if last_content_row is None:
        skipped = [
            item
            for item in skipped
            if not (item.reason == "empty_row" and item.sheet_index == sheet_index)
        ]
    else:
        trimmed: list[FileExtractSkippedRow] = []
        for item in skipped:
            if (
                item.sheet_index == sheet_index
                and item.reason == "empty_row"
                and item.row_index > last_content_row
            ):
                continue
            trimmed.append(item)
        skipped = trimmed

    return items, skipped


def extract_rows_from_file(
    *,
    file: UploadFile,
    content: bytes,
    extract_field_info: str | dict[str, str],
    start_row: int = 2,
    row_count: int = 0,
    sheet_index: int = 1,
) -> FileExtractPublic:
    """
    从 Excel 按字段→列字母映射提取数据。

    - extract_field_info：键=输出字段名，值=列字母（如 A/B）
    - start_row：从第几行开始读（默认 2，表示第 1 行为表头）
    - row_count：计划提取行数；0=从 start_row 起全部；超出实际数据不报错
    - sheet_index=0 读取全部工作表；>=1 为从 1 开始的工作表序号
    - 空行，或任一映射字段为空的行会跳过，并在返回中汇总 sheet/行索引
    """
    mapping = parse_extract_field_info(extract_field_info)
    if start_row < 1:
        raise bad_request("start_row 不能小于 1")
    if row_count < 0:
        raise bad_request("row_count 不能为负数（0=全部）")
    if sheet_index < 0:
        raise bad_request("sheet_index 不能为负数（0=全部，1 起为指定工作表）")

    filename = (file.filename or "").strip().lower()
    if not filename.endswith(_ALLOWED_EXCEL_SUFFIXES):
        raise bad_request("仅支持 .xlsx / .xlsm 格式的 Excel 文件")
    if not content:
        raise bad_request("上传文件为空")

    field_cols = {
        field: column_index_from_string(letter) for field, letter in mapping.items()
    }
    end_row = start_row + row_count - 1 if row_count > 0 else None

    try:
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except InvalidFileException as exc:
        raise bad_request("无法解析 Excel 文件，请确认文件未损坏且为 xlsx/xlsm") from exc
    except Exception as exc:  # noqa: BLE001
        raise bad_request(f"读取 Excel 失败：{exc}") from exc

    try:
        sheets = list(workbook.worksheets)
        if not sheets:
            raise bad_request("Excel 中没有可用的工作表")

        if sheet_index == 0:
            sheet_targets = list(enumerate(sheets, start=1))
        else:
            if sheet_index > len(sheets):
                raise bad_request(
                    f"sheet_index={sheet_index} 超出范围，当前文件共 {len(sheets)} 个工作表"
                )
            sheet_targets = [(sheet_index, sheets[sheet_index - 1])]

        items: list[dict[str, str]] = []
        skipped: list[FileExtractSkippedRow] = []
        for current_sheet_index, sheet in sheet_targets:
            sheet_items, sheet_skipped = _extract_rows_from_sheet(
                sheet,
                sheet_index=current_sheet_index,
                field_cols=field_cols,
                start_row=start_row,
                end_row=end_row,
            )
            items.extend(sheet_items)
            skipped.extend(sheet_skipped)
    finally:
        workbook.close()

    return FileExtractPublic(
        items=items,
        row_count=row_count,
        valid_count=len(items),
        skipped_count=len(skipped),
        skipped_rows=skipped,
    )
